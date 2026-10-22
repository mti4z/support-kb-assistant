"""Client separation: one client must never see another client's documents."""
import io
import uuid

import pytest
from fastapi.testclient import TestClient

from app.db import connect
from app.documents import DocumentError, make_slug, save_document, split_document
from app.main import app
from app.tenants import create_tenant, delete_tenant, get_tenant, hash_key
from app.search import search

client = TestClient(app)

VPN_DOC = (
    "# Acme VPN guide\n\n## Zorblax portal\nTo connect to the zorblax portal open the "
    "Quokka client and enter your badge number.\n\n## Troubleshooting\nIf the quokka "
    "client fails, restart the service desk laptop.\n"
)
OTHER_DOC = (
    "# Beta payroll\n\n## Pagos\nPara cobrar el aguinaldo hay que cargar las horas en el "
    "portal de nómina antes del día 20.\n"
)


def make_pdf(text: str) -> bytes:
    """A tiny valid one-page PDF made by hand (no extra library needed in CI)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)
    return out


@pytest.fixture(scope="module")
def two_tenants():
    a, key_a = create_tenant("Test A " + uuid.uuid4().hex[:6])
    b, key_b = create_tenant("Test B " + uuid.uuid4().hex[:6])
    save_document(a["id"], "vpn-guide.md", VPN_DOC.encode())
    save_document(b["id"], "nomina.md", OTHER_DOC.encode())
    yield a, key_a, b, key_b
    delete_tenant(a["id"])
    delete_tenant(b["id"])


def test_key_is_stored_hashed():
    t, key = create_tenant("Hash test " + uuid.uuid4().hex[:6])
    try:
        with connect() as conn:
            row = conn.execute("SELECT api_key_hash FROM tenants WHERE id = %s", (t["id"],)).fetchone()
        assert row["api_key_hash"] == hash_key(key) and key not in row["api_key_hash"]
        assert get_tenant(key)["id"] == t["id"]
        assert get_tenant("kb_wrong") is None and get_tenant(None) is None
    finally:
        delete_tenant(t["id"])


def test_tenant_finds_own_documents_only(two_tenants):
    a, _, b, _ = two_tenants
    assert search("zorblax quokka badge", tenant_id=a["id"])[0]["slug"] == "vpn-guide"
    assert search("zorblax quokka badge", tenant_id=b["id"]) == []
    assert search("zorblax quokka badge") == []  # the public demo does not see it either


def test_public_search_does_not_leak_tenant_documents(two_tenants):
    assert search("aguinaldo nómina portal") == []
    assert all(p["tenant_id"] is None for p in search("I forgot my password"))


def test_api_ask_uses_the_clients_documents(two_tenants):
    _, key_a, _, key_b = two_tenants
    r = client.post("/ask", json={"question": "how do I use the zorblax quokka badge?"}, headers={"X-API-Key": key_a})
    assert r.status_code == 200 and r.json()["sources"][0]["slug"] == "vpn-guide"
    r = client.post("/ask", json={"question": "how do I use the zorblax quokka badge?"}, headers={"X-API-Key": key_b})
    assert r.json()["mode"] == "no_match" and r.json()["sources"] == []
    r = client.post("/ask", json={"question": "how do I use the zorblax quokka badge?"})
    assert "vpn-guide" not in [x["slug"] for x in r.json()["sources"]]


def test_spanish_client_document(two_tenants):
    _, _, b, key_b = two_tenants
    r = client.post("/ask", json={"question": "¿cómo cobro el aguinaldo?"}, headers={"X-API-Key": key_b})
    assert r.json()["sources"][0]["slug"] == "nomina"


def test_invalid_key_is_rejected_not_ignored():
    for path, call in (("/ask", lambda h: client.post("/ask", json={"question": "password"}, headers=h)),
                       ("/tenant/documents", lambda h: client.get("/tenant/documents", headers=h))):
        assert call({"X-API-Key": "kb_nope"}).status_code == 401
    assert client.get("/tenant/documents").status_code == 401
    assert client.get("/tenant/report").status_code == 401


def test_list_and_delete_documents_are_scoped(two_tenants):
    _, key_a, _, key_b = two_tenants
    docs = client.get("/tenant/documents", headers={"X-API-Key": key_a}).json()["documents"]
    assert [d["slug"] for d in docs] == ["vpn-guide"]
    assert client.delete("/tenant/documents/nomina", headers={"X-API-Key": key_a}).status_code == 404
    assert [d["slug"] for d in client.get("/tenant/documents", headers={"X-API-Key": key_b}).json()["documents"]] == ["nomina"]


def test_same_file_name_in_two_tenants_does_not_collide(two_tenants):
    a, _, b, _ = two_tenants
    save_document(a["id"], "shared-name.txt", b"Alpha plinth manual: the plinth printer plinth room.")
    save_document(b["id"], "shared-name.txt", b"Beta wombat manual: the wombat scanner wombat room.")
    assert search("plinth printer", tenant_id=a["id"])[0]["slug"] == "shared-name"
    assert search("wombat scanner", tenant_id=a["id"]) == []
    assert search("wombat scanner", tenant_id=b["id"])[0]["slug"] == "shared-name"


def test_upload_replaces_same_file(two_tenants):
    a, _, _, _ = two_tenants
    save_document(a["id"], "replace-me.txt", b"The old walrus policy: walrus walrus rules for walrus staff.")
    info = save_document(a["id"], "replace-me.txt", b"The new penguin policy: penguin penguin rules for penguin staff.")
    assert info["replaced"] is True
    assert search("walrus", tenant_id=a["id"]) == []
    assert search("penguin", tenant_id=a["id"])[0]["slug"] == "replace-me"


def test_report_and_feedback_are_scoped(two_tenants):
    _, key_a, _, key_b = two_tenants
    qid = client.post("/ask", json={"question": "something nobody documented"}, headers={"X-API-Key": key_a}).json()["query_id"]
    ra = client.get("/tenant/report", headers={"X-API-Key": key_a}).json()
    rb = client.get("/tenant/report", headers={"X-API-Key": key_b}).json()
    asked = [u["question"] for u in ra["unanswered_questions"]]
    assert "something nobody documented" in asked
    assert "something nobody documented" not in [u["question"] for u in rb["unanswered_questions"]]
    assert client.post("/feedback", json={"query_id": qid, "helpful": True}, headers={"X-API-Key": key_b}).status_code == 404
    assert client.post("/feedback", json={"query_id": qid, "helpful": True}).status_code == 404
    assert client.post("/feedback", json={"query_id": qid, "helpful": True}, headers={"X-API-Key": key_a}).status_code == 200


def test_ingest_keeps_tenant_documents(two_tenants):
    from app.ingest import ingest
    a, _, _, _ = two_tenants
    ingest()
    assert search("zorblax quokka badge", tenant_id=a["id"])[0]["slug"] == "vpn-guide"


def test_upload_endpoint_pdf_docx_txt_and_errors(two_tenants):
    import docx

    _, key_a, _, _ = two_tenants
    d = docx.Document()
    d.add_heading("Flamingo policy", level=1)
    d.add_paragraph("Staff must book the flamingo room: flamingo room rules, flamingo room booking.")
    buf = io.BytesIO()
    d.save(buf)
    files = [
        ("files", ("manual.pdf", make_pdf("Ostrich router reboot: the ostrich router needs a reboot every Monday, reboot ostrich router."), "application/pdf")),
        ("files", ("rooms.docx", buf.getvalue(), "application/octet-stream")),
        ("files", ("notes.txt", b"Llama firewall notes: the llama firewall blocks llama traffic.", "text/plain")),
        ("files", ("bad.exe", b"MZ", "application/octet-stream")),
        ("files", ("empty.txt", b"", "text/plain")),
        ("files", ("broken.pdf", b"not a pdf", "application/pdf")),
    ]
    r = client.post("/tenant/documents", files=files, headers={"X-API-Key": key_a})
    assert r.status_code == 200
    res = {x["file"]: x for x in r.json()["results"]}
    assert res["manual.pdf"]["ok"] and res["rooms.docx"]["ok"] and res["notes.txt"]["ok"]
    assert not res["bad.exe"]["ok"] and not res["empty.txt"]["ok"] and not res["broken.pdf"]["ok"]
    for q, slug in (("ostrich router reboot", "manual"), ("flamingo room", "rooms"), ("llama firewall", "notes")):
        out = client.post("/ask", json={"question": q}, headers={"X-API-Key": key_a}).json()
        assert out["sources"] and out["sources"][0]["slug"] == slug, q


def test_upload_requires_key():
    r = client.post("/tenant/documents", files=[("files", ("a.txt", b"hello world", "text/plain"))])
    assert r.status_code == 401


def test_size_limit(two_tenants):
    a, _, _, _ = two_tenants
    with pytest.raises(DocumentError):
        save_document(a["id"], "big.txt", b"x" * (5 * 1024 * 1024 + 1))


def test_split_plain_text_makes_paragraph_chunks():
    text = "\n\n".join(f"Paragraph {i} " + "word " * 120 for i in range(6))
    title, sections, lang = split_document(text, "long_manual-v2.txt")
    assert title == "long manual v2" and lang == "en"
    assert 3 <= len(sections) <= 6 and all(len(c.split()) <= 260 for _, c in sections)


def test_make_slug():
    assert make_slug("Guía de VPN (v2).pdf") == "guia-de-vpn-v2"
    assert make_slug("C:\\docs\\Mi Manual.docx") == "mi-manual"
    with pytest.raises(DocumentError):
        make_slug("???.txt")
